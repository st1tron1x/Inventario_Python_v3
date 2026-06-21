import sqlite3
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from datetime import datetime


RUTA_DB = "data/inventario.db"



def exportar_excel():

    conexion = sqlite3.connect(RUTA_DB)

    cursor = conexion.cursor()


    cursor.execute("""
    
    SELECT

    c.tipo_documento,
    c.documento,
    c.nombre,
    c.cargo,
    c.area,
    c.sede,
    c.correo,

    e.hostname,
    e.serial,
    e.fabricante,
    e.modelo,
    e.procesador,
    e.memoria,
    e.discos,
    e.ip,
    e.mac,
    e.sistema,
    e.fecha_registro


    FROM equipos e

    INNER JOIN colaboradores c

    ON e.colaborador_id = c.id

    """)


    datos = cursor.fetchall()


    conexion.close()



    wb = Workbook()

    ws = wb.active

    ws.title = "Inventario TI"



    encabezados = [

        "Tipo Documento",
        "Documento",
        "Colaborador",
        "Cargo",
        "Área",
        "Sede",
        "Correo",

        "Hostname",
        "Serial",

        "Fabricante",
        "Modelo",

        "Procesador",

        "Memoria",
        "Discos",

        "IP",
        "MAC",

        "Sistema",

        "Fecha Registro"

    ]


    ws.append(encabezados)


    for celda in ws[1]:

        celda.font = Font(
            bold=True
        )


    for fila in datos:

        ws.append(fila)



    nombre = (
        "Inventario_TYC_"
        +
        datetime.now().strftime(
            "%Y%m%d_%H%M"
        )
        +
        ".xlsx"
    )


    wb.save(nombre)


    return nombre